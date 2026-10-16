from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree


SOURCE = Path(r"D:\works\clients\воздуходувка\Приложение_1.docx")
OUTPUT = Path(r"D:\works\clients\воздуходувка\Приложение_1_с_пояснениями.docx")
MML2OMML = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def set_run_font(run, size=11, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), "Times New Roman")


def format_paragraph(paragraph, size=11, bold=False, italic=False, after=4, before=0):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.08
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, italic=italic)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_font(run, size=14 if level == 1 else 12, bold=True)
    return paragraph


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, size=11, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest, size=11)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=11)
    format_paragraph(paragraph)
    return paragraph


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        border = borders.find(tag)
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:color"), color)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "D9EAF7")
        set_cell_margins(cell)
        if widths:
            cell.width = Cm(widths[index])
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            format_paragraph(paragraph, size=9, bold=True, after=0)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(values):
            cell = cells[column_index]
            cell.text = str(value)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Cm(widths[column_index])
            if row_index % 2:
                set_cell_shading(cell, "F6F9FB")
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index in (0, len(values) - 1) else WD_ALIGN_PARAGRAPH.LEFT
                format_paragraph(paragraph, size=9, after=0)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullet(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.65)
    paragraph.paragraph_format.first_line_indent = Cm(-0.35)
    marker = paragraph.add_run("• ")
    set_run_font(marker, size=11)
    run = paragraph.add_run(text)
    set_run_font(run, size=11)
    format_paragraph(paragraph, after=2)
    return paragraph


def mathml_to_omml(mathml):
    transform = etree.XSLT(etree.parse(str(MML2OMML)))
    result = transform(etree.fromstring(mathml.encode("utf-8")))
    return parse_xml(etree.tostring(result.getroot()))


def add_equation(doc, label, mathml):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(5)
    if label:
        run = paragraph.add_run(label)
        set_run_font(run, size=10, italic=True)
    paragraph._p.append(mathml_to_omml(mathml))
    return paragraph


def replace_math_text(paragraph, old, new, occurrence=1):
    seen = 0
    for node in paragraph._element.xpath(".//m:t"):
        if node.text == old:
            seen += 1
            if seen == occurrence:
                node.text = new
                return True
    return False


doc = Document(SOURCE)

# Preserve native Word equations while correcting the numerical inconsistencies.
paragraphs = doc.paragraphs
replace_math_text(paragraphs[21], "=0,85", "=0,64")
replace_math_text(paragraphs[38], "331", "323")
replace_math_text(paragraphs[38], "≈0,66", "≈0,68")
replace_math_text(paragraphs[39], "=0,86", "=0,88")
replace_math_text(paragraphs[40], "=0,77", "=0,755")
replace_math_text(paragraphs[41], "0,77", "0,755")
replace_math_text(paragraphs[41], "=131590 Па", "=134210 Па")
replace_math_text(paragraphs[42], "0,0404*0,86*131590*0,01", "0,0404*0,88*134210*0,01")
replace_math_text(paragraphs[42], "=2,59 кг/с", "=2,70 кг/с")
replace_math_text(paragraphs[61], "≈035", "≈0,35")

# Make subscripts correspond to the second and third calculation cases.
for paragraph_index in range(36, 43):
    for node in paragraphs[paragraph_index]._element.xpath(".//m:t"):
        if node.text == "1":
            node.text = "2"
replace_math_text(paragraphs[49], "2", "3")
for paragraph_index in range(59, 66):
    for node in paragraphs[paragraph_index]._element.xpath(".//m:t"):
        if node.text == "1":
            node.text = "3"

# Remove an accidental single character while preserving the drawing anchored
# in the same paragraph.
if paragraphs[70].text.strip() == "Ё":
    for text_node in paragraphs[70]._element.xpath(".//w:t"):
        if text_node.text == "Ё":
            text_node.text = ""

# Correct the operational conclusion without overstating an unverified design margin.
paragraphs = doc.paragraphs
for paragraph in paragraphs:
    if paragraph.text.startswith("Массовый расход для компрессора с запасом"):
        paragraph.text = (
            "Максимальный расчётный массовый расход в идеализированной модели составляет 2,70 кг/с "
            "и не содержит эксплуатационного запаса. При коэффициенте запаса 1,10 ориентир составляет "
            "2,97 кг/с. Избыточное полное давление 70 кПа превышает идеальный максимум 56,7 кПа, но "
            "достаточность запаса должна подтверждаться с учётом потерь тракта и совместной рабочей точки G–Δp."
        )
        format_paragraph(paragraph, size=11)
        break

doc.add_page_break()
add_heading(doc, "5 Обозначения и размерности", 1)
add_body(
    doc,
    "Звёздочка у параметра означает полное значение, то есть значение торможения. Индекс i обозначает "
    "расчётный режим или сечение. Давления p и p0 в формулах должны быть абсолютными; избыточное давление "
    "определяется как Δp0 = p0 − pa."
)

symbol_rows = [
    ("ci", "средняя скорость воздуха в рассматриваемом сечении", "м/с"),
    ("pa, pн", "статическое абсолютное давление окружающей среды или в сечении", "Па"),
    ("T, Tн", "статическая термодинамическая температура", "К"),
    ("a, b", "стороны прямоугольного проходного сечения", "м"),
    ("F = ab", "эффективная площадь проходного сечения", "м²"),
    ("cp", "удельная теплоёмкость воздуха при постоянном давлении", "Дж/(кг·К)"),
    ("cv", "удельная теплоёмкость воздуха при постоянном объёме", "Дж/(кг·К)"),
    ("k = cp/cv", "показатель адиабаты", "1"),
    ("R", "удельная газовая постоянная воздуха", "Дж/(кг·К)"),
    ("T0, Ti*", "полная температура потока", "К"),
    ("aкр", "критическая скорость для заданной полной температуры", "м/с"),
    ("λ", "приведённая скорость ci/aкр", "1"),
    ("τ(λ)", "отношение статической температуры к полной T/T0", "1"),
    ("π(λ)", "отношение статического давления к полному p/p0", "1"),
    ("q(λ)", "газодинамическая функция расхода; не динамическое давление", "1"),
    ("m", "размерный коэффициент в формуле расхода", "√К·с/м"),
    ("p0, pi*", "полное абсолютное давление", "Па"),
    ("G, ṁ", "массовый расход воздуха", "кг/с"),
    ("M", "число Маха", "1"),
]
add_table(doc, ("Обозначение", "Физический смысл", "Единица СИ"), symbol_rows, (3.1, 11.3, 3.1))

add_heading(doc, "6 Принятые формулы и допущения", 1)
add_body(
    doc,
    "Расчёт использует изоэнтропические соотношения одномерного течения совершенного газа в форме "
    "газодинамических функций приведённой скорости. Эквивалентные соотношения через число Маха приведены "
    "в материалах NASA Glenn Research Center."
)

MATH_NS = 'xmlns="http://www.w3.org/1998/Math/MathML"'
add_equation(doc, "Полная температура: ", f'<math {MATH_NS}><mrow><msub><mi>T</mi><mn>0</mn></msub><mo>=</mo><mi>T</mi><mo>+</mo><mfrac><msup><mi>c</mi><mn>2</mn></msup><mrow><mn>2</mn><msub><mi>c</mi><mi>p</mi></msub></mrow></mfrac></mrow></math>')
add_equation(doc, "Критическая скорость: ", f'<math {MATH_NS}><mrow><msub><mi>a</mi><mtext>кр</mtext></msub><mo>=</mo><msqrt><mrow><mfrac><mrow><mn>2</mn><mi>k</mi></mrow><mrow><mi>k</mi><mo>+</mo><mn>1</mn></mrow></mfrac><mi>R</mi><msub><mi>T</mi><mn>0</mn></msub></mrow></msqrt></mrow></math>')
add_equation(doc, "Приведённая скорость: ", f'<math {MATH_NS}><mrow><mi>λ</mi><mo>=</mo><mfrac><mi>c</mi><msub><mi>a</mi><mtext>кр</mtext></msub></mfrac></mrow></math>')
add_equation(doc, "Температурная функция: ", f'<math {MATH_NS}><mrow><mi>τ</mi><mo>(</mo><mi>λ</mi><mo>)</mo><mo>=</mo><mn>1</mn><mo>−</mo><mfrac><mrow><mi>k</mi><mo>−</mo><mn>1</mn></mrow><mrow><mi>k</mi><mo>+</mo><mn>1</mn></mrow></mfrac><msup><mi>λ</mi><mn>2</mn></msup></mrow></math>')
add_equation(doc, "Функция давления: ", f'<math {MATH_NS}><mrow><mi>π</mi><mo>(</mo><mi>λ</mi><mo>)</mo><mo>=</mo><msup><mrow><mo>[</mo><mi>τ</mi><mo>(</mo><mi>λ</mi><mo>)</mo><mo>]</mo></mrow><mfrac><mi>k</mi><mrow><mi>k</mi><mo>−</mo><mn>1</mn></mrow></mfrac></msup></mrow></math>')
add_equation(doc, "Функция расхода: ", f'<math {MATH_NS}><mrow><mi>q</mi><mo>(</mo><mi>λ</mi><mo>)</mo><mo>=</mo><mi>λ</mi><msup><mrow><mo>(</mo><mfrac><mrow><mi>k</mi><mo>+</mo><mn>1</mn></mrow><mn>2</mn></mfrac><mo>)</mo></mrow><mfrac><mn>1</mn><mrow><mi>k</mi><mo>−</mo><mn>1</mn></mrow></mfrac></msup><msup><mrow><mo>[</mo><mi>τ</mi><mo>(</mo><mi>λ</mi><mo>)</mo><mo>]</mo></mrow><mfrac><mn>1</mn><mrow><mi>k</mi><mo>−</mo><mn>1</mn></mrow></mfrac></msup></mrow></math>')
add_equation(doc, "Коэффициент расхода: ", f'<math {MATH_NS}><mrow><mi>m</mi><mo>=</mo><msqrt><mfrac><mi>k</mi><mi>R</mi></mfrac></msqrt><msup><mrow><mo>(</mo><mfrac><mn>2</mn><mrow><mi>k</mi><mo>+</mo><mn>1</mn></mrow></mfrac><mo>)</mo></mrow><mfrac><mrow><mi>k</mi><mo>+</mo><mn>1</mn></mrow><mrow><mn>2</mn><mo>(</mo><mi>k</mi><mo>−</mo><mn>1</mn><mo>)</mo></mrow></mfrac></msup></mrow></math>')
add_equation(doc, "Массовый расход: ", f'<math {MATH_NS}><mrow><mi>G</mi><mo>=</mo><mfrac><mrow><mi>m</mi><mi>q</mi><mo>(</mo><mi>λ</mi><mo>)</mo><msub><mi>p</mi><mn>0</mn></msub><mi>F</mi></mrow><msqrt><msub><mi>T</mi><mn>0</mn></msub></msqrt></mfrac></mrow></math>')

add_body(doc, "Допущения расчёта:", bold_lead="Допущения расчёта:")
for item in (
    "сухой воздух рассматривается как совершенный газ; k, R и cp приняты постоянными;",
    "течение стационарное, адиабатическое и изоэнтропическое; потери полного давления до расчётного сечения не учитываются;",
    "скорость, давление и температура равномерны по эффективной площади сечения;",
    "pн = 101 325 Па является статическим абсолютным давлением в расчётном сечении;",
    "полученные значения описывают идеализированный поток, а не готовую рабочую точку компрессора и сети."
):
    add_bullet(doc, item)
add_body(
    doc,
    "Значения cp = 1006 Дж/(кг·К), k = 1,4 и R = 287 Дж/(кг·К) близки, но не строго согласованы: "
    "из kR/(k−1) следует cp ≈ 1004,5 Дж/(кг·К). Для данного уровня точности влияние невелико, однако в "
    "окончательном расчёте следует использовать один согласованный набор свойств при заданной температуре."
)

add_heading(doc, "7 Проверка расчётных случаев", 1)
check_rows = [
    ("1", "0,0075", "280", "327,116", "0,8460", "0,6411", "0,9714", "158,06", "2,573"),
    ("2", "0,0100", "220", "312,206", "0,6804", "0,7550", "0,8781", "134,21", "2,696"),
    ("3", "0,0150", "110", "294,164", "0,3505", "0,9302", "0,5250", "108,93", "2,022"),
]
add_table(
    doc,
    ("Режим", "F, м²", "c, м/с", "T0, К", "λ", "π", "q", "p0, кПа", "G, кг/с"),
    check_rows,
    (1.4, 1.8, 1.9, 2.0, 1.5, 1.5, 1.5, 2.0, 2.1),
)
add_body(
    doc,
    "Независимая проверка по G = ρcF при ρ = p/(RT) даёт 2,573; 2,695 и 2,022 кг/с. Это подтверждает "
    "пересчёт. Во втором режиме исходный документ содержал ошибочный знаменатель 331 вместо 323 м/с, "
    "из-за чего расход был занижен с 2,70 до 2,59 кг/с. В первом режиме подпись π(λ) исправлена с 0,85 "
    "на 0,64; само последующее вычисление давления уже использовало 0,64."
)
add_body(
    doc,
    "Для выбора оборудования требуется не менее 2,70 кг/с без запаса. При условном коэффициенте 1,10 "
    "получается 2,97 кг/с. Требуемое полное избыточное давление первого режима равно примерно 56,7 кПа; "
    "70 кПа соответствует запасу около 23 %, пока не учтены потери фильтра, воздуховодов, решётки, сопла и "
    "местные сопротивления. Расход и давление необходимо проверять совместно по характеристике компрессора и сети."
)

add_heading(doc, "8 Критика расчёта боковых сил", 1)
add_body(
    doc,
    "Контуры скорости и давления иллюстрируют структуру течения, но сами по себе не подтверждают силы "
    "13,3, 88 и 77 кгс. Для воспроизводимого инженерного вывода нужен отчёт интегральной силы по выбранным "
    "стенкам объекта, включая давление и вязкое трение, с направлением осей и знаком компоненты. Основной "
    "единицей следует использовать ньютон; кгс можно оставить справочно, 1 кгс = 9,80665 Н."
)
add_body(doc, "В текущей версии не указаны:", bold_lead="В текущей версии не указаны:")
for item in (
    "трёхмерная или двумерная постановка, толщина модели и точное расстояние 100 мм;",
    "граничные условия, массовый расход или полное давление на входе, модель плотности и включение уравнения энергии;",
    "модель турбулентности, качество сетки, y+, сеточная сходимость и чувствительность к размеру расчётной области;",
    "невязки, баланс массы и импульса, история сходимости силы и опорное давление;",
    "поверхности интегрирования, компоненты Fx, Fy, Fz и разделение силы давления и вязкого трения;",
    "для нестационарного решения — средняя, среднеквадратичная и пиковая сила."
):
    add_bullet(doc, item)
add_body(
    doc,
    "Оценка масштаба по потоку импульса Gc даёт примерно 721, 593 и 222 Н для трёх режимов. Поэтому значения "
    "88 и 77 кгс выше полного осевого импульса соответствующих струй примерно в 1,20 и 1,27 раза. Это возможно "
    "при дополнительной составляющей давления, развороте потока или влиянии границ, но требует контрольного "
    "баланса. Значение 13,3 кгс составляет около 0,59Gс и выглядит правдоподобно при неполном перехвате струи."
)
add_body(
    doc,
    "Запас 100 кгс относительно 88 кгс составляет только 13,6 % и пока не учитывает неопределённость CFD и "
    "динамические пики. До принятия этого усилия как расчётного следует провести три расчёта на последовательно "
    "сгущаемых сетках, тест размера области, контроль баланса массы и импульса, мониторинг силы и хотя бы одну "
    "экспериментальную проверку тензодатчиком."
)

add_heading(doc, "9 Что добавить в расчётный отчёт", 1)
for item in (
    "Таблицу режимов с F, c, G, p0, Δp0, расстоянием до объекта, Fx, Fy, Fz, Gc и отношением F/(Gc).",
    "Схему с размерами сопла, объекта, расчётной области, направлением осей и поверхностями интегрирования силы.",
    "Таблицу граничных условий и моделей: compressible/ideal gas, energy, turbulence, wall treatment, operating pressure.",
    "Отчёт по сеточной независимости, балансу массы и импульса, невязкам и истории сходимости интегральной силы.",
    "Характеристику компрессора и сети на одном графике, включая потери тракта и условия воздуха.",
    "Оценку потребной мощности привода и допустимых температур, шума, вибрации и пульсаций потока.",
    "Расчёт направляющих по силе в ньютонах с коэффициентом запаса, трением, динамическими пиками и ресурсом."
):
    add_bullet(doc, item)

add_heading(doc, "10 Источники формул и методов проверки", 1)
sources = [
    ("NASA Glenn Research Center. Isentropic Flow Equations", "https://www.grc.nasa.gov/www/k-12/airplane/isentrop.html"),
    ("NASA Glenn Research Center. Mass Flow Rate Equations", "https://www1.grc.nasa.gov/beginners-guide-to-aeronautics/mass-flow-rate-equations/"),
    ("NASA Glenn Research Center. Derivation of the Isentropic Flow Equations", "https://www.grc.nasa.gov/WWW/k-12/VirtualAero/BottleRocket/airplane/isndrv508.html"),
    ("Ansys Fluent Theory Guide. Computing Forces, Moments, and the Center of Pressure", "https://ansyshelp.ansys.com/public/Views/Secured/corp/v242/en/flu_th/flu_th_sec_report_force_moment.html"),
    ("NASA Standard 7009B. Standard for Models and Simulations", "https://standards.nasa.gov/standard/nasa/nasa-std-7009"),
    ("NASA. Examining Spatial Grid Convergence", "https://ntrs.nasa.gov/api/citations/20080014300/downloads/20080014300.pdf?attachment=true"),
    ("NIST Special Publication 330. Definition of the standard atmosphere", "https://www.nist.gov/pml/special-publication-330/sp-330-appendix-1"),
]
for source_index, (title, url) in enumerate(sources, start=1):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.65)
    paragraph.paragraph_format.first_line_indent = Cm(-0.65)
    number_run = paragraph.add_run(f"{source_index}. ")
    set_run_font(number_run, size=10)
    title_run = paragraph.add_run(title + ". ")
    set_run_font(title_run, size=10)
    url_run = paragraph.add_run(url)
    set_run_font(url_run, size=9)
    format_paragraph(paragraph, size=10, after=3)

for section in doc.sections:
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)

doc.save(OUTPUT)
print(OUTPUT)
