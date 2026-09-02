"""Build the RecorderLnx user guide DOCX from the maintained Markdown source."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from PIL import Image


ROOT = Path(r"D:\works\OburecGH\Lazarus\RecorderLnx\Docs\Руководство пользователя")
SOURCE = ROOT / "Руководство пользователя RecorderLnx.md"
OUTPUT = ROOT / "Руководство пользователя RecorderLnx.docx"
BLUE = "2E74B5"
LIGHT_BLUE = "E8EEF5"
GRID = "AAB7C4"


def set_font(run, name="Calibri", size=11, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Страница ")
    set_font(run, size=9, color="667788")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def configure_section(section, landscape=False):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    if landscape:
        section.page_width, section.page_height = Mm(279.4), Mm(215.9)
        section.left_margin = section.right_margin = Inches(0.65)
    else:
        section.page_width, section.page_height = Mm(215.9), Mm(279.4)
        section.left_margin = section.right_margin = Inches(1.0)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    header = section.header.paragraphs[0]
    header.text = "RecorderLnx — руководство пользователя"
    set_font(header.runs[0], size=9, color="667788")
    add_page_field(section.footer.paragraphs[0])


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, before, after, color in (
        ("Heading 1", 16, 18, 10, BLUE),
        ("Heading 2", 13, 14, 7, BLUE),
        ("Heading 3", 12, 10, 5, "1F4D78"),
        ("Heading 4", 11, 8, 4, "1F4D78"),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_inline_markdown(paragraph, text):
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, name="Consolas", size=9.5, color="1F4D78")
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, bold=True)
        else:
            paragraph.add_run(part)


def clean_cell(text):
    return text.strip().strip("|").strip()


def add_table(doc, rows):
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    widths = [0.72, 1.35, 3.75, 0.68] if cols == 4 else [6.5 / cols] * cols
    if doc.sections[-1].orientation == WD_ORIENT.LANDSCAPE:
        widths = [0.8, 1.8, 6.0, 1.0] if cols == 4 else [9.6 / cols] * cols
    for r_idx, values in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.width = Inches(widths[c_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_inline_markdown(p, values[c_idx] if c_idx < len(values) else "")
            for run in p.runs:
                if r_idx == 0:
                    run.bold = True
                run.font.size = Pt(8.5 if cols == 4 else 9)
        if r_idx == 0:
            repeat_header(table.rows[r_idx])
        prevent_row_split(table.rows[r_idx])
    return table


def add_image(doc, image_path, caption):
    with Image.open(image_path) as image:
        width, height = image.size
    max_width = 9.6 if doc.sections[-1].orientation == WD_ORIENT.LANDSCAPE else 6.5
    max_height = 5.45 if doc.sections[-1].orientation == WD_ORIENT.LANDSCAPE else 6.45
    display_width = min(max_width, max_height * width / height)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(image_path), width=Inches(display_width))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption)
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(8)
    run = cp.add_run(caption)
    set_font(run, size=9, color="667788")


def build():
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0], False)
    doc.core_properties.title = "RecorderLnx — руководство пользователя"
    doc.core_properties.subject = "Эксплуатация и настройка RecorderLnx"
    doc.core_properties.author = "MERA"

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    landscape_active = False
    first_title = True
    while i < len(lines):
        line = lines[i].rstrip()
        if not line or line == "---":
            i += 1
            continue
        image_match = re.match(r"!\[([^]]*)\]\(([^)]+)\)", line)
        heading_match = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading_match:
            next_image_wide = False
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            if j < len(lines):
                next_match = re.match(r"!\[([^]]*)\]\(([^)]+)\)", lines[j])
                if next_match:
                    with Image.open(ROOT / next_match.group(2)) as next_image:
                        next_image_wide = next_image.width / max(next_image.height, 1) > 1.3
            if landscape_active and not next_image_wide:
                section = doc.add_section(WD_SECTION.NEW_PAGE)
                configure_section(section, False)
                landscape_active = False
            elif next_image_wide and not landscape_active:
                section = doc.add_section(WD_SECTION.NEW_PAGE)
                configure_section(section, True)
                landscape_active = True
            level = len(heading_match.group(1))
            title = heading_match.group(2)
            if first_title:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(120)
                p.paragraph_format.space_after = Pt(12)
                run = p.add_run(title)
                set_font(run, size=30, bold=True, color="1F4D78")
                first_title = False
            else:
                doc.add_heading(title, level=min(level, 4))
            i += 1
            continue
        if image_match:
            path = ROOT / image_match.group(2)
            with Image.open(path) as image:
                wide = image.width / max(image.height, 1) > 1.3
            if wide and not landscape_active:
                section = doc.add_section(WD_SECTION.NEW_PAGE)
                configure_section(section, True)
                landscape_active = True
            add_image(doc, path, image_match.group(1))
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            rows = [[clean_cell(c) for c in line.strip().strip("|").split("|")]]
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([clean_cell(c) for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, rows)
            continue
        if re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            add_inline_markdown(p, re.sub(r"^\d+\.\s+", "", line))
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_markdown(p, line[2:])
            i += 1
            continue
        if line.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            set_cell_shading_like = OxmlElement("w:shd")
            set_cell_shading_like.set(qn("w:fill"), "F2F4F7")
            p._p.get_or_add_pPr().append(set_cell_shading_like)
            run = p.add_run("\n".join(block))
            set_font(run, name="Consolas", size=8.5)
            continue
        paragraph_lines = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,4})\s+|^!\[|^\||^- |^\d+\.\s+|^```|^---$", lines[i]):
            paragraph_lines.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        add_inline_markdown(p, " ".join(paragraph_lines))

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
