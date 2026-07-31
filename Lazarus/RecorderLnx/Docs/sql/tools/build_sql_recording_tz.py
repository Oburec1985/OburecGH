from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "sql-recording-module-tz.md"
OUTPUT = ROOT / "sql-recording-module-tz.docx"
ACCENT = RGBColor(31, 78, 121)
MUTED = RGBColor(95, 105, 115)


def font(run, size=10.5, bold=False, italic=False, color=None):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    font(run, 9, color=MUTED)


def add_rich_text(paragraph, text, size=10.5, italic=False, color=None):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            font(run, size - 0.3, color=color)
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            font(run, size, bold=True, italic=italic, color=color)
        else:
            run = paragraph.add_run(part)
            font(run, size, italic=italic, color=color)


def new_numbering_instance(doc):
    numbering = doc.part.numbering_part.element
    style_num_id = int(doc.styles["List Number"]._element.pPr.numPr.numId.val)
    base_num = next(n for n in numbering.findall(qn("w:num")) if int(n.get(qn("w:numId"))) == style_num_id)
    abstract_id = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    next_id = max([int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))] + [0]) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(next_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return next_id


def add_list_item(doc, text, numbered=False, num_id=None):
    style = "List Number" if numbered else "List Bullet"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Inches(0.38)
    p.paragraph_format.first_line_indent = Inches(-0.19)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if numbered and num_id is not None:
        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = num_id
    add_rich_text(p, text)


def add_table(doc, rows):
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(1.8), Inches(4.7)] if cols == 2 else [Inches(6.5 / cols)] * cols
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.width = widths[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_rich_text(p, value, size=9.1)
            if r_idx == 0:
                set_cell_shading(cell, "D9E7F5")
                for run in p.runs:
                    run.bold = True
    set_repeat_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    doc.settings.odd_and_even_pages_header_footer = False

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, before, after in (("Heading 1", 15, 14, 7), ("Heading 2", 12.5, 10, 5), ("Heading 3", 11, 8, 4)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(header.add_run("RecorderLnx  |  ТЗ модуля SQL-записи"), 8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run("Страница "), 9, color=MUTED)
    add_page_number(footer)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(5)
    font(title.add_run("ТЕХНИЧЕСКОЕ ЗАДАНИЕ"), 22, bold=True, color=ACCENT)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    font(subtitle.add_run("Модуль SQL-записи проекта RecorderLnx"), 15, bold=True)
    for label, value in (("Статус", "Концепция / требования для согласования"), ("Версия", "0.1"), ("Дата", "31.07.2026"), ("Очередь", "Настройка и runtime-запись")):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        font(p.add_run(label + ": "), 10, bold=True, color=MUTED)
        font(p.add_run(value), 10)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(12)
    p_pr = rule._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "14"), ("space", "1"), ("color", "1F4E79")):
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    p_pr.append(borders)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 5  # title and metadata are already rendered as a masthead
    active_num_id = None
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("| "):
            active_num_id = None
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                values = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[-: ]+", x) for x in values):
                    rows.append(values)
                i += 1
            if rows:
                add_table(doc, rows)
            continue
        image_match = re.match(r"!\[(.+?)\]\((.+?)\)", line)
        if image_match:
            active_num_id = None
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            path = ROOT / image_match.group(2)
            max_width = 6.55
            p.add_run().add_picture(str(path), width=Inches(max_width))
            i += 1
            continue
        if line.startswith("### "):
            active_num_id = None
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            active_num_id = None
            if line.startswith("## Приложение"):
                doc.add_page_break()
            doc.add_heading(line[3:], level=1)
        elif re.match(r"^\d+\. ", line):
            if active_num_id is None:
                active_num_id = new_numbering_instance(doc)
            add_list_item(doc, re.sub(r"^\d+\. ", "", line), numbered=True, num_id=active_num_id)
        elif line.startswith("- "):
            active_num_id = None
            add_list_item(doc, line[2:])
        elif line.startswith("*") and line.endswith("*"):
            active_num_id = None
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(8)
            add_rich_text(p, line.strip("*"), size=9, italic=True, color=MUTED)
        elif line.startswith("**"):
            active_num_id = None
            p = doc.add_paragraph()
            add_rich_text(p, line)
        else:
            active_num_id = None
            p = doc.add_paragraph()
            p.paragraph_format.keep_together = False
            add_rich_text(p, line)
        i += 1

    props = doc.core_properties
    props.title = "Техническое задание: модуль SQL-записи RecorderLnx"
    props.subject = "Настройка и runtime-запись данных мониторинга"
    props.author = "Проект RecorderLnx"
    props.keywords = "RecorderLnx, SQL, база данных, запись, события, файлы"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
