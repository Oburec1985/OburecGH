from __future__ import annotations

import re
import zipfile
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


SOURCE_ROOT = Path(r"D:\works\OburecGH\Lazarus\RecorderLnx")
OUTPUT_DOCX = Path(r"D:\works\OburecGH\docs\Rlnx\RecorderLnx_all_sources.docx")

PRIMARY_SOURCE_SUFFIXES = {
    ".pas",
    ".lfm",
    ".lpr",
    ".lpi",
    ".lpg",
    ".inc",
    ".rc",
}

SCRIPT_SUFFIXES = {
    "",
    ".bat",
    ".cmd",
    ".ps1",
    ".py",
    ".sh",
    ".desktop",
}

EXCLUDED_DIRS = {
    ".git",
    "_buildverify",
    "backup",
    "cach",
    "docs",
    "lib",
    "logs",
    "rendered",
    "tmp",
    "tests",
}

SECTION_TITLES = {
    ".": "Корень проекта",
    "calibrations": "Калибровки",
    "core": "Ядро и модель проекта",
    "device": "Устройства и сбор данных",
    "export": "Экспорт и внешние форматы",
    "icons": "Иконки и ресурсы интерфейса",
    "linuxsetupmanager": "LinuxSetupManager",
    "pages": "Страницы и визуальные компоненты",
    "plugins": "Плагины",
    "sdb": "База характеристик",
    "scripts": "Скрипты",
    "storage": "Хранение и запись",
    "tools": "Инструменты",
    "ui": "Пользовательский интерфейс",
    "utils": "Вспомогательные модули",
}


@dataclass(frozen=True)
class SourceFile:
    path: Path
    rel: str
    section_key: str
    section_title: str
    encoding: str
    text: str
    line_count: int
    char_count: int


def should_skip(path: Path) -> bool:
    rel_parts = [part.lower() for part in path.relative_to(SOURCE_ROOT).parts]
    return any(part in EXCLUDED_DIRS for part in rel_parts[:-1])


def is_script_path(path: Path) -> bool:
    rel_parts = [part.lower() for part in path.relative_to(SOURCE_ROOT).parts]
    return "scripts" in rel_parts or "tools" in rel_parts


def is_source_file(path: Path) -> bool:
    suffix = path.suffix.lower()
    if suffix in PRIMARY_SOURCE_SUFFIXES:
        return True
    if is_script_path(path) and suffix in SCRIPT_SUFFIXES:
        return True
    return False


def iter_candidate_files() -> Iterable[Path]:
    for path in SOURCE_ROOT.rglob("*"):
        if path.is_file() and not should_skip(path) and is_source_file(path):
            yield path


def read_text(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1"), "latin-1"


def sanitize_for_docx(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\x00", "")

    def is_valid_xml_char(char: str) -> bool:
        code = ord(char)
        return (
            code in (0x09, 0x0A, 0x0D)
            or 0x20 <= code <= 0xD7FF
            or 0xE000 <= code <= 0xFFFD
        )

    return "".join(char if is_valid_xml_char(char) else " " for char in text)


def section_for(path: Path) -> tuple[str, str]:
    rel = path.relative_to(SOURCE_ROOT)
    if len(rel.parts) == 1:
        key = "."
    else:
        key = rel.parts[0].lower()
    return key, SECTION_TITLES.get(key, rel.parts[0])


def collect_sources() -> list[SourceFile]:
    sources: list[SourceFile] = []
    for path in sorted(iter_candidate_files(), key=lambda item: str(item.relative_to(SOURCE_ROOT)).lower()):
        raw_text, encoding = read_text(path)
        text = sanitize_for_docx(raw_text)
        rel = str(path.relative_to(SOURCE_ROOT))
        section_key, section_title = section_for(path)
        sources.append(
            SourceFile(
                path=path,
                rel=rel,
                section_key=section_key,
                section_title=section_title,
                encoding=encoding,
                text=text,
                line_count=text.count("\n") + (1 if text else 0),
                char_count=len(text),
            )
        )
    return sources


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(8)
    run.font.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.1)
    section.bottom_margin = Cm(1.1)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    styles["Normal"].font.size = Pt(9)

    for style_name in ("Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.color.rgb = RGBColor(0, 0, 0)

    code_style = styles.add_style("SourceCode", 1)
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code_style._element.rPr.rFonts.set(qn("w:cs"), "Consolas")
    code_style.font.size = Pt(6.8)
    code_style.paragraph_format.space_before = Pt(0)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.line_spacing = 1.0

    meta_style = styles.add_style("FileMeta", 1)
    meta_style.font.name = "Calibri"
    meta_style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    meta_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    meta_style.font.size = Pt(8)
    meta_style.font.italic = True
    meta_style.paragraph_format.space_after = Pt(3)
    return doc


def add_title(doc: Document, sources: list[SourceFile]) -> None:
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("RecorderLnx полный листинг исходников")

    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    paragraph = doc.add_paragraph()
    paragraph.add_run(
        f"Автоматически сгенерировано {now}. В документ включены текстовые исходники "
        f"RecorderLnx и служебные скрипты из {SOURCE_ROOT}. Исключены каталоги Docs, "
        "Tests, cach, tmp, lib и сборочные артефакты."
    )

    total_lines = sum(item.line_count for item in sources)
    total_chars = sum(item.char_count for item in sources)
    paragraph = doc.add_paragraph()
    paragraph.add_run(
        f"Всего файлов: {len(sources)}. Строк: {total_lines}. Символов: {total_chars}."
    )


def add_manifest(doc: Document, sources: list[SourceFile]) -> None:
    doc.add_heading("Реестр файлов", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ("Раздел", "Файл", "Строк", "Символов", "Кодировка")
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True)
        set_cell_shading(table.rows[0].cells[index], "D9EAF7")

    for item in sources:
        cells = table.add_row().cells
        set_cell_text(cells[0], item.section_title)
        set_cell_text(cells[1], item.rel)
        set_cell_text(cells[2], str(item.line_count))
        set_cell_text(cells[3], str(item.char_count))
        set_cell_text(cells[4], item.encoding)


def iter_line_chunks(text: str, lines_per_chunk: int = 90) -> Iterable[str]:
    lines = text.split("\n")
    for index in range(0, len(lines), lines_per_chunk):
        yield "\n".join(lines[index : index + lines_per_chunk])


def add_sources(doc: Document, sources: list[SourceFile]) -> None:
    doc.add_page_break()
    doc.add_heading("Полный текст исходников", level=1)
    by_section: dict[str, list[SourceFile]] = defaultdict(list)
    for item in sources:
        by_section[item.section_key].append(item)

    for section_key in sorted(by_section, key=lambda key: SECTION_TITLES.get(key, key)):
        section_sources = by_section[section_key]
        doc.add_heading(section_sources[0].section_title, level=1)
        for item in section_sources:
            doc.add_heading(item.rel, level=2)
            doc.add_paragraph(
                f"Путь: {item.path} | Кодировка: {item.encoding} | Строк: {item.line_count}",
                style="FileMeta",
            )
            if item.text:
                for chunk in iter_line_chunks(item.text):
                    paragraph = doc.add_paragraph(style="SourceCode")
                    run = paragraph.add_run(chunk)
                    run.font.name = "Consolas"
                    run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
                    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
                    run._element.rPr.rFonts.set(qn("w:cs"), "Consolas")
                    run.font.size = Pt(6.8)
            else:
                doc.add_paragraph("[Пустой файл]", style="SourceCode")


def build_document() -> None:
    sources = collect_sources()
    doc = setup_document()
    add_title(doc, sources)
    add_manifest(doc, sources)
    add_sources(doc, sources)
    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)

    with zipfile.ZipFile(OUTPUT_DOCX) as package:
        xml = package.read("word/document.xml").decode("utf-8")
    replacement_sources = [
        item.rel for item in sources if "\ufffd" in item.text
    ]

    print(f"Saved {OUTPUT_DOCX}")
    print(f"Included files: {len(sources)}")
    print(f"Included lines: {sum(item.line_count for item in sources)}")
    print(f"Included chars: {sum(item.char_count for item in sources)}")
    if replacement_sources:
        print("Warning: replacement characters already exist in these source files:")
        for rel in replacement_sources:
            print(f"  - {rel}")
    elif "\ufffd" in xml:
        raise RuntimeError("DOCX contains replacement characters not present in source text")


if __name__ == "__main__":
    build_document()
