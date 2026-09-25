from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


DOCX_PATH = Path(r"D:\works\OburecGH\docs\Rlnx\Презентация Rlnx\Сценарий презентации Rlnx.docx")
FORM_IMAGE_PATH = Path(r"D:\works\OburecGH\docs\screen\recorder\Формуляр.png")
TEMP_PATH = DOCX_PATH.with_name(DOCX_PATH.stem + ".updated.docx")


def insert_paragraph_before(reference: Paragraph, text: str, style: str) -> Paragraph:
    paragraph_xml = OxmlElement("w:p")
    reference._p.addprevious(paragraph_xml)
    paragraph = Paragraph(paragraph_xml, reference._parent)
    paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def delete_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


document = Document(DOCX_PATH)

# The first figure is the mnemonic/formular example. Replacing the image-part
# blob preserves the established size, position and caption formatting.
first_image_paragraph = next(
    paragraph
    for paragraph in document.paragraphs
    if paragraph._p.xpath(".//a:blip")
)
first_blip = first_image_paragraph._p.xpath(".//a:blip")[0]
first_image_rid = first_blip.get(
    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
)
document.part.related_parts[first_image_rid]._blob = FORM_IMAGE_PATH.read_bytes()

# Remove the narrow industry-specific STT/VKU chapter. The presentation is
# intended to describe the product in general terms.
paragraphs = list(document.paragraphs)
industry_start = next(
    index
    for index, paragraph in enumerate(paragraphs)
    if paragraph.text.startswith("12.")
)
industry_end = next(
    index
    for index, paragraph in enumerate(paragraphs[industry_start + 1 :], industry_start + 1)
    if paragraph.text.startswith("13.")
)
for paragraph in paragraphs[industry_start:industry_end]:
    delete_paragraph(paragraph)

# Renumber the following general section after removing the industry chapter.
next_heading = next(
    paragraph for paragraph in document.paragraphs if paragraph.text.startswith("13.")
)
next_heading.text = "12." + next_heading.text[3:]
next_heading.style = "Heading 1"

# Remove remaining explicit STT/VKU bullets or sentences without touching
# general descriptions of diagnostics and automation.
for paragraph in list(document.paragraphs):
    text = paragraph.text
    if "СТТ" in text or "ВКУ" in text:
        if paragraph.style.name == "Feature bullet":
            delete_paragraph(paragraph)
        else:
            paragraph.text = text.replace("СТТ и ВКУ", "прикладных измерительных систем")
            paragraph.text = paragraph.text.replace("СТТ/ВКУ", "прикладных измерительных систем")

# Add a compact product-level block before the vector-channel figure. It
# explains capabilities, not the settings procedure.
vector_image_paragraph = list(
    paragraph for paragraph in document.paragraphs if paragraph._p.xpath(".//a:blip")
)[1]
insert_paragraph_before(vector_image_paragraph, "Уставки и управление", "Heading 2")
setpoint_bullets = [
    "•  четыре независимые границы для каждого тега: верхняя и нижняя предупредительные, верхняя и нижняя аварийные;",
    "•  индивидуальные пороги, включение, сообщения, цвета и гистерезис срабатывания;",
    "•  единое состояние уставок для индикаторов, мнемосхем, SVG-компонентов, журнала событий и внешней логики;",
    "•  чтение и изменение порогов, а также включение и отключение уставок из Lua-сценариев во время работы системы;",
    "•  автоматическая реакция сценария на состояние тега: изменение представления, расчётов и управляющих действий проекта.",
]
for bullet in setpoint_bullets:
    insert_paragraph_before(vector_image_paragraph, bullet, "Feature bullet")

# Make the Lua capability explicit in the plugin section as well.
plugin_heading_index = next(
    index
    for index, paragraph in enumerate(document.paragraphs)
    if paragraph.text.startswith("7.")
)
sql_heading = next(
    paragraph
    for paragraph in document.paragraphs[plugin_heading_index + 1 :]
    if paragraph.text.startswith("8.")
)
insert_paragraph_before(
    sql_heading,
    "•  Lua API для чтения и изменения уставок тегов непосредственно в пользовательском сценарии;",
    "Feature bullet",
)

document.save(TEMP_PATH)
TEMP_PATH.replace(DOCX_PATH)

