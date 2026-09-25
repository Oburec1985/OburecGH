from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(r"D:\works\OburecGH")
SOURCE = ROOT / r"docs\Rlnx\РП на ПО\БЛИЖ.409801.100.236-01 34_v03.docx"
OUTPUT = ROOT / r"docs\Rlnx\РП на ПО\БЛИЖ.409801.100.236-01 34_v04.docx"
SCREENS = ROOT / r"Lazarus\RecorderLnx\Docs\Руководство пользователя\screens"
UA = SCREENS / "user-annotated"


def find_paragraph(doc, prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise RuntimeError(f"Anchor not found: {prefix}")


def insert_before(anchor, element):
    anchor._p.addprevious(element)


def make_paragraph(doc, text="", style=None, align=None, bold=False, keep_next=False):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.font.name = "Times New Roman"
        r.font.size = Pt(14)
    if align is not None:
        p.alignment = align
    if keep_next:
        p.paragraph_format.keep_with_next = True
    return p


def cloned_heading(doc, text):
    p_xml = deepcopy(HEADING_TEMPLATE._p)
    for child in list(p_xml):
        if child.tag != qn("w:pPr"):
            p_xml.remove(child)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    bold = OxmlElement("w:b")
    rpr.append(bold)
    run.append(rpr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    p_xml.append(run)
    return p_xml


def add_border_table(doc, rows):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    # Keep item ranges and element names readable without tab stops or
    # character-by-character wrapping in Word's narrow table cells.
    widths = [0.90, 1.60, 2.10, 1.50]
    headers = ["№", "Элемент", "Как работает", "Характерные действия"]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = text
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "D9EAF7")
        cell._tc.get_or_add_tcPr().append(shd)
    for num, element, behavior, action in rows:
        cells = table.add_row().cells
        for i, text in enumerate((num, element, behavior, action)):
            cells[i].text = text
            cells[i].width = Inches(widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9)
            if i == 0:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                # Number ranges are identifiers, not prose: Word must not split
                # them across lines even when the neighboring columns are busy.
                cells[i]._tc.get_or_add_tcPr().append(OxmlElement("w:noWrap"))
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    return table


def add_figure_block(doc, anchor, heading, intro, image_path, caption, rows, width=6.1):
    # The preceding tag-properties table already fills its page.  An explicit
    # break before the next block would therefore create a header-only page.
    if not heading.startswith("4.1.15 "):
        page = make_paragraph(doc)
        page.add_run().add_break(WD_BREAK.PAGE)
        insert_before(anchor, page._p)

    insert_before(anchor, cloned_heading(doc, heading))

    p = make_paragraph(doc, intro)
    p.paragraph_format.keep_with_next = True
    insert_before(anchor, p._p)

    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.keep_with_next = True
    pic.add_run().add_picture(str(image_path), width=Inches(width))
    insert_before(anchor, pic._p)

    cap = make_paragraph(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
    for run in cap.runs:
        run.italic = True
        run.font.size = Pt(11)
    insert_before(anchor, cap._p)

    table = add_border_table(doc, rows)
    insert_before(anchor, table._tbl)


doc = Document(SOURCE)
HEADING_TEMPLATE = find_paragraph(doc, "4.1.10 Диагностика и сохранность данных")
anchor_42 = find_paragraph(doc, "4.2 Модуль RCPanel")

figures = [
    (
        "4.1.11 Основное окно и диагностика RecorderLnx",
        "Основное окно объединяет формуляры проекта, команды режима работы, состояние записи, поиск тегов и журнал. Перед переходом в Record оператор проверяет состояние источников и отсутствие критических сообщений.",
        SCREENS / "01-main-form.png",
        "Рисунок 1 — Основное окно RecorderLnx и средства оперативной диагностики",
        [
            ("1–2", "Заголовок и страницы", "Показывают открытый проект и доступные формуляры.", "Выбрать требуемую страницу проекта."),
            ("3", "Команды Stop, Preview, Record", "Изменяют состояние измерительного цикла и файловой записи.", "Сначала проверить данные в Preview, затем включить Record."),
            ("4–5", "Поиск и список тегов", "Находят тег по имени и показывают его текущее состояние.", "Использовать для быстрой проверки канала."),
            ("6–10", "Журнал и фильтры", "Показывают сообщения системы, данных, тревог и устройств.", "При диагностике включать только нужные категории."),
        ],
    ),
    (
        "4.1.12 Общие параметры проекта и записи",
        "Общие настройки задаются до измерения в режиме Stop. Здесь определяются каталоги, параметры обновления, предыстория и автоматические условия начала и завершения записи.",
        SCREENS / "02-recorder-settings.png",
        "Рисунок 2 — Окно общих настроек RecorderLnx",
        [
            ("1–4", "Вкладки и обновление", "Разделяют настройки регистратора, оборудования, каналов и плагинов; задают темп обновления интерфейса.", "Не путать с частотой опроса прибора."),
            ("5–6", "Испытание и изделие", "Формируют идентификацию сеанса и метаданные записи.", "Заполнить до начала Record."),
            ("7–19", "Запись и каталоги", "Задают предысторию, паузы, сохранение конфигурации, рабочий каталог и базу ГХ.", "Проверить путь и свободное место."),
            ("20–35", "Условия старта и останова", "Запускают и завершают запись по команде, уровню, триггеру или времени.", "Порог задавать в конечных единицах тега."),
            ("36–38", "Применение настроек", "Сохраняют либо отменяют изменения.", "После применения проверить проект в Preview."),
        ],
    ),
    (
        "4.1.13 Источники данных и поиск устройств",
        "В аппаратной конфигурации оператор формирует состав источников, задаёт адреса и проверяет связь. Автопоиск не добавляет прибор без подтверждения пользователя.",
        UA / "user-config-hardware.png",
        "Рисунок 3 — Настройка источников данных и проверка связи",
        [
            ("1–3", "Дерево и свойства источника", "Показывают состав проекта и параметры выбранного прибора.", "Выбрать источник перед редактированием."),
            ("4–7", "Интерфейс, адрес и порт", "Определяют канал связи с устройством.", "Изменять в Stop; затем выполнить проверку TCP."),
            ("8–11", "Добавление и удаление", "Меняют состав источников проекта.", "Удаление применять только после проверки зависимых тегов."),
            ("12–15", "Автопоиск и диагностика", "Ищут доступные приборы и проверяют сетевую доступность.", "Добавить отмеченные устройства и проверить их в Preview."),
        ],
    ),
    (
        "4.1.14 Свойства тега и цепочка преобразований",
        "Тег является единой точкой данных для формуляров, расчётов, уставок, файлов и SQL. Единицы и диапазон должны соответствовать выходу конечной градуировочной характеристики.",
        UA / "user-tag-main.png",
        "Рисунок 4 — Основные свойства тега и цепочка градуировочных характеристик",
        [
            ("1–7", "Имя, адрес и описание", "Связывают тег с физическим либо виртуальным каналом.", "Имя должно быть уникальным; адрес проверяется по источнику."),
            ("8–10", "Диапазон", "Задаёт ожидаемые границы физической величины и масштаб компонентов.", "Проверить в Preview до Record."),
            ("11–12", "Аппаратная ГХ", "Применяет преобразование, связанное с каналом прибора.", "Не изменять из мнемосхемы."),
            ("13–14", "Канальные ГХ", "Последовательно преобразуют результат аппаратной ГХ.", "Контролировать порядок характеристик и единицы."),
            ("15–17", "Подтверждение", "Сохраняет или отменяет изменения тега.", "После сохранения проверить уставки и оси графиков."),
        ],
        4.9,
    ),
    (
        "4.1.15 Выбор градуировочной характеристики",
        "Утверждённые характеристики выбираются из базы SDB/Mera Files. Перед применением проверяются метаданные, единицы, диапазон и форма зависимости.",
        UA / "user-calibration-db.png",
        "Рисунок 5 — Выбор градуировочной характеристики из базы",
        [
            ("1–3", "Дерево базы", "Группирует доступные характеристики и позволяет выбрать запись.", "Использовать утверждённую запись нужного датчика."),
            ("4–7", "Метаданные", "Показывают имя, описание, единицы и идентификаторы характеристики.", "Сверить с паспортом канала."),
            ("8–10", "График и диапазон", "Показывают зависимость входа и выхода.", "Проверить монотонность и рабочую область."),
            ("11–12", "Выбор и отмена", "Копируют выбранную ГХ в проект либо закрывают диалог.", "После выбора проверить показания в Preview."),
        ],
    ),
    (
        "4.1.16 Формуляры и мнемосхемы",
        "Формуляр представляет пользовательскую страницу отображения. Он хранит компоновку и привязки компонентов, но не заменяет теги и не перепрограммирует прибор.",
        UA / "user-base-page.png",
        "Рисунок 6 — Рабочая страница RecorderLnx с мнемосхемой",
        [
            ("1–2", "Вкладки страниц", "Переключают формуляры проекта и позволяют работать с несколькими представлениями.", "Для отдельного монитора страницу можно отцепить."),
            ("3–6", "Компоненты мнемосхемы", "Отображают значения, графики и состояние объекта по привязанным тегам.", "Настройка компонента меняет только представление."),
            ("7–9", "Команды режима", "Управляют Stop, Preview и Record.", "Перед Record проверить единицы, диапазоны и тревоги."),
            ("10–12", "SQL и журнал", "Включают дополнительную запись оценок и показывают сообщения.", "SQL не заменяет основной файл сырых данных."),
        ],
    ),
    (
        "4.1.17 Осциллограмма и оперативный тренд",
        "Осциллограмма предназначена для анализа формы быстрого сигнала, а оперативный тренд — для наблюдения оценок текущего сеанса. Масштабирование изменяет только отображение.",
        UA / "user-oscillogram.png",
        "Рисунок 7 — Осциллограмма блока измерительных данных",
        [
            ("1–3", "Панель и легенда", "Выбирают линии, показывают значения и состояние курсоров.", "Сопоставить цвет линии с осью Y."),
            ("4–6", "Область сигнала", "Показывает отсчёты векторного тега во временном окне.", "Мышью выбрать интересующий участок."),
            ("7–9", "Оси и масштаб", "Задают видимый диапазон времени и амплитуды.", "Автомасштаб не изменяет исходные данные."),
            ("10–12", "Триггер и курсоры", "Фиксируют кадр по условию и измеряют интервалы/уровни.", "Проверить источник и уровень триггера."),
        ],
    ),
    (
        "4.1.18 Спектральное представление",
        "Спектр строится по векторному тегу. Частотный диапазон и разрешение определяются частотой дискретизации, длиной блока и настройками окна.",
        UA / "user-spectrum.png",
        "Рисунок 8 — Спектр и оценки по частотным полосам",
        [
            ("1–3", "Линии и легенда", "Показывают спектры выбранных каналов и их текущие оценки.", "Сопоставить линию с нужной осью."),
            ("4–7", "Частотный график", "Отображает амплитуду по частоте, курсоры и выбранные полосы.", "Проверить единицы и тип амплитуды."),
            ("8–10", "Параметры анализа", "Задают БПФ, окно, усреднение и интегрирование.", "Подбирать по требуемому разрешению."),
            ("11–13", "Полосовые оценки", "Вычисляют значения в заданных диапазонах частот.", "Использовать утверждённые границы полос."),
        ],
    ),
    (
        "4.1.19 Настройка SQL базы данных",
        "SQL-подсистема сохраняет выбранные оценки и метаданные в Firebird. Она включается отдельно и не заменяет основную файловую запись.",
        SCREENS / "12-sql-db-settings.png",
        "Рисунок 9 — Настройка записи измерений в SQL БД",
        [
            ("1–13", "Подключение и защита", "Задают backend, файл/сервер БД, порт, пользователя, пароль или переменную окружения.", "Для удалённой БД указывать адрес сервера, не localhost."),
            ("14–16", "Период и очередь", "Определяют частоту точек SQL и запас отложенных записей.", "Не путать с частотой опроса прибора."),
            ("17–21", "Объект мониторинга", "Записывают имя, тип и серийный номер объекта.", "Заполнить до рабочего сеанса."),
            ("22–28", "Теги и оценки", "Выбирают каналы, производные величины и управляющий тег.", "Назначить только необходимые оценки."),
            ("29–31", "Проверка и применение", "Проверяют Firebird, сервис и схему, затем сохраняют настройки.", "Проверку выполнить до Record."),
        ],
    ),
    (
        "4.1.20 Исторический SQL-тренд",
        "SQL-тренд читает сохранённые точки независимо от текущего измерительного блока. Ограничение числа отображаемых точек не удаляет данные.",
        UA / "user-sql-trend.png",
        "Рисунок 10 — Просмотр исторических данных SQL",
        [
            ("1–3", "Отображения и линии", "Группируют исторические каналы, подписи и цвета.", "Назначить каждую линию нужной оси Y."),
            ("4–7", "График и курсор", "Показывают архивный интервал и значение в выбранной точке.", "Масштабировать по времени и активной оси."),
            ("8–10", "Интервал времени", "Выбирают фиксированный диапазон либо скользящее окно текущей даты.", "Проверить часовой пояс и границы."),
            ("11–13", "Обновление и удаление", "Перечитывают БД; удаление интервала является отдельной необратимой операцией.", "Перед удалением обязательно создать резервную копию БД."),
        ],
    ),
]

for item in figures:
    add_figure_block(doc, anchor_42, *item)

# RCPanel: the application is a single main form with four tabs. A compact
# operator-facing table is included even when a clean repository screenshot is
# unavailable; HostAgent has no separate GUI.
anchor_43 = find_paragraph(doc, "4.3 Модуль HostAgent")
insert_before(anchor_43, cloned_heading(doc, "4.2.1 Диалоги RCPanel"))
p = make_paragraph(doc, "RCPanel содержит одно главное окно с четырьмя вкладками. Состояние HostAgent и команды удалённого запуска показываются на вкладке хостов; отдельного интерфейса HostAgent нет.")
insert_before(anchor_43, p._p)
rc_rows = [
    ("1", "Вкладка «RecorderLnx хосты»", "Показывает связь с компьютером, HostAgent и состояние Rlnx как независимые признаки.", "Выбрать хост; выполнить Preview, Record, Stop, запуск, Wake-on-LAN или разрешённое выключение."),
    ("2", "Вкладка «События записи»", "Показывает сгруппированные сеансы и связанные пакеты измерений.", "Задать период, обновить, открыть пакеты либо удалить только запись события из SQL."),
    ("3", "Диалог «MERA-пакеты события»", "Показывает файлы, хост, состояние, размер и путь каждого пакета.", "Открыть, перенести, исправить путь или сохранить описание события."),
    ("4", "Вкладка «Файловые хранилища»", "Задаёт локальный или SMB-каталог; SFTP в текущей реализации не используется как рабочий транспорт.", "Добавить хранилище и проверить доступность каталога."),
    ("5", "Вкладка «Журнал»", "Показывает результат команд, обнаружения, SQL, SDB и HostAgent.", "Использовать для диагностики; записи журнала не являются измерительными данными."),
]
table = add_border_table(doc, rc_rows)
insert_before(anchor_43, table._tbl)

# Normalize every table, including the inherited v03 tables: literal tabs and
# tab-stop definitions are removed, and body text is never justified. This
# prevents Word from producing tab-like stretched spaces inside cells.
for table in doc.tables:
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            for paragraph in cell.paragraphs:
                # Table text must not inherit body-paragraph indents: in narrow
                # cells they look like tabs and force identifiers to wrap.
                paragraph.paragraph_format.left_indent = Pt(0)
                paragraph.paragraph_format.right_indent = Pt(0)
                paragraph.paragraph_format.first_line_indent = Pt(0)
                for run in paragraph.runs:
                    if "\t" in run.text:
                        run.text = run.text.replace("\t", " ")
                    if col_index == 0 and "–" in run.text:
                        run.text = run.text.replace("–", "‑")
                    if col_index == 0:
                        run.font.size = Pt(8)
                ppr = paragraph._p.get_or_add_pPr()
                tabs = ppr.find(qn("w:tabs"))
                if tabs is not None:
                    ppr.remove(tabs)
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                    if row_index == 0 or col_index == 0
                    else WD_ALIGN_PARAGRAPH.LEFT
                )

# Ask Word to refresh TOC and fields on open.
settings = doc.settings._element
update = settings.find(qn("w:updateFields"))
if update is None:
    update = OxmlElement("w:updateFields")
    settings.append(update)
update.set(qn("w:val"), "true")

doc.save(OUTPUT)
print(OUTPUT)
