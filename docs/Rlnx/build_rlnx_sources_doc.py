from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


REPO_ROOT = Path(r"D:\works\OburecGH")
PROJECT_ROOT = REPO_ROOT / "Lazarus" / "RecorderLnx"
OUTPUT_DIR = REPO_ROOT / "docs" / "Rlnx"
OUTPUT_DOCX = OUTPUT_DIR / "RecorderLnx_source_guide.docx"

SOURCE_SUFFIXES = {".pas", ".lpr", ".lfm", ".lpi", ".lpg", ".inc"}
EXCLUDED_PARTS = {"Docs", "Tests", "cach", "_buildverify", "lib", "backup"}


SECTION_RULES = [
    ("Запуск и проект", ("RecorderLnx.lpr", "RecorderLnx.lpi", "group.lpg", "Plugins/plugins.lpg")),
    ("Ядро и жизненный цикл", ("Core/",)),
    ("Алгоритмы", ("Algs/",)),
    ("Устройства и сбор данных", ("Device/",)),
    ("Пользовательский интерфейс", ("UI/",)),
    ("Настройки и системные инструменты", ("Tools/LinuxSetupManager/",)),
    ("Калибровки и SDB", ("Calibrations/", "SDB/")),
    ("SQL и архив", ("SQLdb/",)),
    ("Плагины", ("Plugins/",)),
    ("Корневые формы и совместимость", ("u",)),
]

SECTION_INTRO = {
    "Запуск и проект": (
        "Точка входа, Lazarus-проекты и группы проектов задают состав собираемых целей. "
        "Эти файлы важны не столько алгоритмами, сколько связями между модулями, ресурсами и целями сборки."
    ),
    "Ядро и жизненный цикл": (
        "Core хранит переносимую модель RecorderLnx: теги, состояние записи, конфигурацию, очереди событий, "
        "MERA-файлы, пути ресурсов и сетевой клиент координатора. UI должен обращаться к этим модулям как к доменной модели, "
        "а не хранить собственные копии состояния."
    ),
    "Алгоритмы": (
        "Раздел Algs содержит вычислительные блоки: спектр, полосы частот и runtime-обвязку алгоритмов. "
        "Эти модули должны оставаться отделёнными от форм, потому что работают многократно во время измерений."
    ),
    "Устройства и сбор данных": (
        "Device объединяет общий контракт источников данных и реализации конкретных устройств MIC-140, MIC-185 и MCbus. "
        "Главная идея раздела - явный жизненный цикл: поиск, подключение, конфигурация, запуск сбора, остановка и отключение."
    ),
    "Пользовательский интерфейс": (
        "UI содержит LCL-формы, визуальные компоненты, страницы мнемосхем, графики, диалоги настроек и контроллер редактора форм. "
        "Большая часть процедур здесь связывает пользовательские действия с Core-моделью."
    ),
    "Настройки и системные инструменты": (
        "LinuxSetupManager и родственные файлы автоматизируют системные настройки Linux-стендов: сеть, SSH, Samba, время, доступ, диски, профили и ассоциации файлов."
    ),
    "Калибровки и SDB": (
        "Калибровки и SDB отвечают за хранение градуировочных характеристик, выбор характеристик из базы, свойства SDB-дерева и тензо-калибровки."
    ),
    "SQL и архив": (
        "SQLdb содержит настройки подключения, репозиторий, runtime-запись, импорт выбора каналов, исторический SQL-тренд и события MERA-архива."
    ),
    "Плагины": (
        "Плагинная подсистема пока начинается с ABI метаданных: SDK, образец DLL/SO и чтение PLUGININFO без создания runtime-плагина."
    ),
    "Корневые формы и совместимость": (
        "В корне остаются совместимые копии LFM и отдельные единицы, которые исторически нужны проекту и Lazarus-ресурсам. "
        "Их лучше рассматривать как слой совместимости, а не как место для новой доменной логики."
    ),
}

KEY_FILES = {
    "RecorderLnx.lpr": "Главная точка входа приложения: инициализация LCL, single-instance guard, старт основной формы и базовых сервисов.",
    "UI/uMainForm.pas": "Главный экран, размещение компонентов, команды тулбара и связка UI с контроллером приложения.",
    "Core/uRecorderApplicationController.pas": "Координирует жизненный цикл приложения и связывает Core-сервисы с действиями верхнего уровня.",
    "Core/uRecorderCoreServices.pas": "Собирает набор доменных сервисов, которыми пользуются формы и runtime-логика.",
    "Core/uRecorderTags.pas": "Модель тегов, их свойства, привязки к источникам данных и runtime-значения.",
    "Core/uRecorderDataSources.pas": "Абстракции источников данных и поток передачи измерений в теги.",
    "Core/uRecorderDataStorage.pas": "Запись данных в MERA/хранилище и управление каналами записи.",
    "Core/uRecorderStateMachine.pas": "Состояния RecorderLnx и допустимые переходы между stop, preview, record и служебными режимами.",
    "Core/uRecorderEventQueue.pas": "Очередь событий между runtime-слоем и UI, чтобы не смешивать потоки и визуальные обновления.",
    "Core/uMeraFile.pas": "Низкоуровневая работа с MERA-файлами и структурой записанного пакета.",
    "Device/uRecorderDeviceInterfaces.pas": "Базовые интерфейсы устройств и источников; граница между Core и конкретной аппаратурой.",
    "Device/uRecorderDeviceDataThread.pas": "Общий шаблон потокового сбора данных с управлением стадиями и ошибками.",
    "Device/MIC140/uRecorderMic140DataSource.pas": "Runtime-источник MIC-140: конфигурация каналов, чтение блоков и публикация данных.",
    "Device/MIC140/uRecorderMic140Protocol.pas": "Протокольные команды MIC-140 и разбор сетевого обмена.",
    "Device/mic185/uRecorderMic185DataSource.pas": "Runtime-источник MIC-185 с каналами, группами и аппаратным временем.",
    "Device/mic185/uMic185MebiusTcpProtocol.pas": "TCP-протокол Mebius для MIC-185.",
    "Algs/uRecorderSpectrumEngine.pas": "Численное ядро спектральных расчётов.",
    "Algs/uRecorderSpectrumRuntime.pas": "Runtime-обвязка спектра и взаимодействие с тегами/буферами.",
    "UI/uRecorderTrendView.pas": "Визуальный тренд live-данных.",
    "UI/uRecorderOglOscillogramView.pas": "OpenGL-осциллограмма для быстрого отображения сигналов.",
    "UI/uRecorderMeasurementSectionView.pas": "Визуальный компонент измерительного сечения и расчёта механических величин.",
    "UI/uRecorderSettingsDialog.pas": "Главный диалог настроек приложения, оборудования, SQL, путей и системных параметров.",
    "UI/uTagSettingsDialog.pas": "Диалог настройки тегов и их привязки к источникам/ГХ/записи.",
    "SQLdb/uRecorderSqlDbRepository.pas": "Доступ к SQL-схеме, запросам и данным архива.",
    "SQLdb/SqlTrend/uRecorderSqlTrendView.pas": "Отображение исторических данных и событий из SQL.",
    "SDB/uRecorderSdbStore.pas": "Хранилище SDB-дерева и файлов градуировочных характеристик.",
    "Calibrations/Strains/uRecorderStrainCalibration.pas": "Модель тензо-калибровки и формулы преобразования.",
    "Tools/LinuxSetupManager/uLinuxSetupManagerMain.pas": "Главная форма LinuxSetupManager.",
    "Plugins/SDK/uRecorderPluginApi.pas": "Общий SDK ABI для плагинов RecorderLnx.",
}

ARCHITECTURE_NOTES = [
    (
        "Core first pipeline",
        "Основной поток данных выглядит как устройства или файлы -> IRecorderDataSource -> TRecorderTagRegistry и буферы тегов -> EventBus и snapshot-очереди -> UI, запись, SQL и алгоритмы. Такая схема удерживает формы от прямой работы с аппаратурой и потоками.",
    ),
    (
        "Единый фасад TRecorder",
        "Корневой объект ядра собирает шину событий, теги, источники данных, машину состояний, время, спектр, тревоги и SQL. UI должен работать через этот фасад и сервисы, а не напрямую с конкретными MIC или MCbus-классами.",
    ),
    (
        "Жизненный цикл записи",
        "Состояния Stop, Preview и Record обслуживаются state machine и application controller. Переход состояния подготавливает алгоритмы, запускает время, сбор данных и запись, а обратный переход закрывает runtime-ресурсы.",
    ),
    (
        "Потоки и UI",
        "Подписчики EventBus вызываются в потоке публикации, поэтому для LCL используются snapshot-очереди. Worker-потоки не должны трогать визуальные контролы напрямую.",
    ),
    (
        "Runtime источников",
        "Контракт источника сводится к стадиям ConfigureTags, PrepareHardware, Start, Tick или DoTick, Stop. Конкретные MIC-140, MIC-185 и MCbus скрыты за общими интерфейсами.",
    ),
    (
        "Запись и SQL",
        "MERA writer пишет блоки тегов в файловый пакет, а SQL runtime ставит задания в writer-thread. Соединения, транзакции и репозиторий SQL живут вне GUI-потока.",
    ),
]

UI_NOTES = [
    (
        "Главная форма",
        "TMainForm является центром сценариев: открыть проект, настроить теги и оборудование, включить предпросмотр, начать запись, сохранить результат, открыть WinPOS и управлять SQL-записью.",
    ),
    (
        "Формуляры и мнемосхемы",
        "Формуляр - пользовательская страница проекта. TRecorderFormPage хранит компоненты, фон, режим страницы и параметры откреплённого окна; редактор размещает компоненты вручную.",
    ),
    (
        "Палитра компонентов",
        "Палитра строится из фабрик компонентов. Группа 'Графики' объединяет осциллограмму, тренд, SQL-тренд и спектр; MainForm больше не должен вручную знать каждый тип.",
    ),
    (
        "Графики",
        "Осциллограмма, тренд и спектр являются разными способами смотреть одни и те же теги: мгновенную форму сигнала, историю во времени и частотный анализ.",
    ),
    (
        "Настройки проекта",
        "Главный диалог настроек управляет общей конфигурацией, источниками, каналами, виртуальными тегами, алгоритмами и плагинами. Он меняет проектную модель и runtime-реестр, а не только внешний вид.",
    ),
    (
        "Приборные диалоги",
        "Диалоги MIC-140, MIC-185 и MCbus описывают сценарий подключения источника: найти или проверить связь, выбрать каналы, задать аппаратные режимы и связать каналы с тегами.",
    ),
    (
        "Измерительное сечение",
        "Компонент измерительного сечения является прикладным сценарием: назначает тензотеги ролям, рассчитывает механические величины, показывает недопустимые данные и поддерживает пакетный импорт таблиц.",
    ),
]


@dataclass
class SourceFile:
    rel: str
    section: str
    suffix: str
    lines: int
    classes: list[str]
    routines: list[str]
    comment_lines: int
    has_intro_comment: bool
    role: str
    comment: str
    snippet: str


def read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "cp1251", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def section_for(rel: str) -> str:
    rel_norm = rel.replace("\\", "/")
    for section, prefixes in SECTION_RULES:
        for prefix in prefixes:
            if prefix == "u":
                if "/" not in rel_norm and rel_norm.lower().startswith("u"):
                    return section
            elif rel_norm == prefix.rstrip("/") or rel_norm.startswith(prefix):
                return section
    return "Прочие исходники"


def extract_decl(pattern: str, text: str, limit: int) -> list[str]:
    found = []
    for match in re.finditer(pattern, text, flags=re.IGNORECASE | re.MULTILINE):
        name = match.group(1)
        if name not in found:
            found.append(name)
        if len(found) >= limit:
            break
    return found


def has_intro_comment(lines: list[str]) -> bool:
    checked = 0
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        checked += 1
        if stripped.startswith(("//", "{", "(*")):
            return True
        if checked >= 8:
            return False
    return False


def role_for(rel: str, section: str, classes: list[str]) -> str:
    if rel in KEY_FILES:
        return KEY_FILES[rel]
    name = Path(rel).stem
    lower = rel.lower()
    if rel.endswith(".lfm"):
        return "Lazarus-описание формы или frame: хранит состав визуальных контролов, размеры, подписи и привязку обработчиков."
    if rel.endswith(".lpi") or rel.endswith(".lpg"):
        return "Файл проекта или группы Lazarus: задаёт состав модулей, параметры компиляции и связи целей."
    if "settingsdialog" in lower:
        return "Диалог настроек: читает текущую модель, показывает параметры пользователю и применяет изменения обратно."
    if "datasource" in lower:
        return "Runtime-источник данных: превращает настройки оборудования в поток измерений и публикацию значений тегов."
    if "protocol" in lower:
        return "Протокольный слой: описывает команды, структуры пакетов и правила обмена с внешним устройством или сервисом."
    if "calibration" in lower or "sdb" in lower:
        return "Калибровочный или SDB-модуль: хранение, выбор и применение градуировочных характеристик."
    if "trend" in lower or "oscillogram" in lower or "spectrum" in lower:
        return "Модуль отображения или расчёта графиков: подготовка данных и визуальная навигация по сигналам."
    if "thread" in lower or "runtime" in lower:
        return "Runtime-модуль: выполняется во время работы приложения и должен избегать тяжёлой логики в GUI-потоке."
    if classes:
        return f"Модуль объявляет {', '.join(classes[:3])}; роль определяется связями раздела {section}."
    return f"Служебный исходник раздела {section}; назначение уточняется по месту использования."


def comment_for(source: SourceFile | None, rel: str, section: str, has_comment: bool) -> str:
    if has_comment:
        return "В исходнике уже есть комментарии или явные имена; в документе фиксируется роль и границы ответственности."
    if section == "Ядро и жизненный цикл":
        return "Комментарий: держать этот модуль независимым от LCL-форм; UI должен вызывать его через сервисы и события."
    if section == "Устройства и сбор данных":
        return "Комментарий: различать стадии поиска, подключения, программирования, сбора и отключения; не смешивать их в обработчиках UI."
    if section == "Пользовательский интерфейс":
        return "Комментарий: форма должна связывать действия пользователя с моделью, но не становиться источником доменных данных."
    if section == "Алгоритмы":
        return "Комментарий: этот код выполняется многократно во время работы; важны предсказуемые буферы, отсутствие лишних выделений и явные единицы измерения."
    if section == "SQL и архив":
        return "Комментарий: SQL-модуль должен явно различать локальные пути, сетевые URI, схему БД и runtime-состояние записи."
    if section == "Калибровки и SDB":
        return "Комментарий: не смешивать ссылочные характеристики из SDB и локальные копии, изменённые внутри проекта."
    if section == "Настройки и системные инструменты":
        return "Комментарий: системные действия должны иметь предпросмотр, понятную диагностику и отдельный безопасный CLI-путь."
    return "Комментарий: назначение модуля стоит держать в одном предложении рядом с объявлением основных типов."


def make_snippet(text: str) -> str:
    lines = text.splitlines()
    candidates = []
    for i, line in enumerate(lines):
        if re.search(r"^\s*(type|interface|implementation)\b", line, re.IGNORECASE):
            candidates.append(i)
        if re.search(r"^\s*T[A-Za-z0-9_]+\s*=\s*class|^\s*function\s+|^\s*procedure\s+", line, re.IGNORECASE):
            candidates.append(i)
        if len(candidates) >= 2:
            break
    if not candidates:
        start = 0
    else:
        start = max(0, candidates[0] - 2)
    snippet = "\n".join(lines[start:start + 18])
    snippet = re.sub(r"\t", "  ", snippet)
    return snippet.strip()[:1800]


def collect_sources() -> list[SourceFile]:
    items: list[SourceFile] = []
    for path in PROJECT_ROOT.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in SOURCE_SUFFIXES:
            continue
        rel_path = path.relative_to(PROJECT_ROOT)
        if any(part in EXCLUDED_PARTS for part in rel_path.parts):
            continue
        rel = rel_path.as_posix()
        text = read_text(path)
        lines = text.splitlines()
        classes = extract_decl(r"^\s*(T[A-Za-z0-9_]+)\s*=\s*class\b", text, 8)
        classes += [x for x in extract_decl(r"^\s*(I[A-Za-z0-9_]+)\s*=\s*interface\b", text, 4) if x not in classes]
        routines = extract_decl(r"^\s*(?:class\s+)?(?:procedure|function)\s+([A-Za-z0-9_.]+)", text, 10)
        comment_lines = sum(1 for line in lines if line.strip().startswith(("//", "{", "(*", "*")))
        section = section_for(rel)
        intro = has_intro_comment(lines)
        role = role_for(rel, section, classes)
        item = SourceFile(
            rel=rel,
            section=section,
            suffix=path.suffix.lower(),
            lines=len(lines),
            classes=classes[:8],
            routines=routines[:10],
            comment_lines=comment_lines,
            has_intro_comment=intro,
            role=role,
            comment="",
            snippet=make_snippet(text),
        )
        item.comment = comment_for(item, rel, section, intro)
        items.append(item)
    return sorted(items, key=lambda x: (x.section, x.rel.lower()))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(8.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table, widths: list[float] | None = None) -> None:
    table.style = "Table Grid"
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            if row_idx == 0:
                set_cell_shading(cell, "1F4E79")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.color.rgb = RGBColor(255, 255, 255)
                        r.bold = True
            if widths and col_idx < len(widths):
                cell.width = Cm(widths[col_idx])


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True)
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            set_cell_text(cells[i], value)
    style_table(table, widths)


def add_code_block(document: Document, text: str) -> None:
    if not text:
        return
    p = document.add_paragraph()
    p.style = "Code"
    r = p.add_run(text)
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    r.font.size = Pt(7.5)


def ensure_styles(document: Document) -> None:
    styles = document.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
    if "Code" not in [s.name for s in styles]:
        style = styles.add_style("Code", 1)
        style.font.name = "Consolas"
        style.font.size = Pt(7.5)
        style.paragraph_format.space_before = Pt(2)
        style.paragraph_format.space_after = Pt(6)


def build_doc() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    sources = collect_sources()
    by_section: dict[str, list[SourceFile]] = {}
    for item in sources:
        by_section.setdefault(item.section, []).append(item)

    doc = Document()
    ensure_styles(doc)
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.4)

    title = doc.add_paragraph(style="Title")
    title.add_run("Исходники RecorderLnx").bold = True
    subtitle = doc.add_paragraph()
    subtitle.add_run("Структурированный обзор, карта модулей и аннотированный указатель исходного кода").italic = True

    doc.add_paragraph(
        "Документ собран как рабочая карта исходников RecorderLnx. Он группирует файлы по подсистемам, "
        "показывает назначение ключевых модулей и добавляет текстовые комментарии там, где роль файла не очевидна из заголовка. "
        "Цель документа - быстро понять, где искать алгоритмы, UI, настройки, устройства, SQL-архив и вспомогательные инструменты."
    )
    doc.add_paragraph(
        f"Источник: {PROJECT_ROOT}. В индекс включены {len(sources)} основных файлов .pas/.lfm/.lpr/.lpi/.lpg/.inc без Docs, Tests, cach, _buildverify и бинарных артефактов."
    )

    doc.add_heading("Общая карта подсистем", level=1)
    summary_rows = []
    for section_name in [rule[0] for rule in SECTION_RULES] + ["Прочие исходники"]:
        items = by_section.get(section_name, [])
        if not items:
            continue
        pas_count = sum(1 for x in items if x.suffix == ".pas")
        lfm_count = sum(1 for x in items if x.suffix == ".lfm")
        summary_rows.append([
            section_name,
            str(len(items)),
            f"Pascal: {pas_count}; формы: {lfm_count}",
            SECTION_INTRO.get(section_name, "Дополнительные исходники и связующие файлы проекта."),
        ])
    add_table(doc, ["Раздел", "Файлов", "Состав", "Назначение"], summary_rows, [4.0, 1.4, 3.2, 9.0])

    doc.add_heading("Как читать этот документ", level=1)
    for text in [
        "Сначала используйте разделы 2-10 как карту подсистем; затем переходите к таблице файлов внутри нужного раздела.",
        "Поле 'Комментарий' не заменяет комментарии в исходнике. Это внешняя поясняющая заметка к роли файла и рискам сопровождения.",
        "Для файлов без вводного комментария документ добавляет пояснение о границе ответственности: что модуль должен делать и чего лучше не смешивать в нём.",
        "Кодовые фрагменты приведены только для ключевых файлов, чтобы документ оставался навигационным, а не превращался в нечитабельный полный листинг.",
    ]:
        doc.add_paragraph(text, style=None)

    doc.add_heading("Архитектурный обзор Core и runtime", level=1)
    doc.add_paragraph(
        "Этот раздел суммирует карту ядра, устройств и алгоритмов. Он помогает читать последующие таблицы не как список файлов, "
        "а как цепочку исполнения приложения."
    )
    add_table(doc, ["Тема", "Комментарий"], [[a, b] for a, b in ARCHITECTURE_NOTES], [4.5, 12.5])

    doc.add_heading("Обзор UI и пользовательских сценариев", level=1)
    doc.add_paragraph(
        "UI-слой лучше понимать через сценарии оператора и редактора проекта: настройка источников, настройка тегов, размещение компонентов, "
        "предпросмотр, запись, графики и архив."
    )
    add_table(doc, ["Сценарий", "Комментарий"], [[a, b] for a, b in UI_NOTES], [4.5, 12.5])

    key_by_section = {}
    for rel in KEY_FILES:
        for item in sources:
            if item.rel == rel:
                key_by_section.setdefault(item.section, []).append(item)
                break

    for section_name in [rule[0] for rule in SECTION_RULES] + ["Прочие исходники"]:
        items = by_section.get(section_name, [])
        if not items:
            continue
        doc.add_heading(section_name, level=1)
        doc.add_paragraph(SECTION_INTRO.get(section_name, "Дополнительная группа исходников проекта."))
        rows = []
        for item in items:
            entities = ", ".join(item.classes[:4]) or ", ".join(item.routines[:3]) or "-"
            rows.append([
                item.rel,
                str(item.lines),
                entities,
                item.role,
                item.comment,
            ])
        add_table(doc, ["Файл", "Строк", "Ключевые сущности", "Назначение", "Комментарий"], rows, [4.2, 1.2, 3.1, 5.3, 5.0])

        keys = key_by_section.get(section_name, [])
        if keys:
            doc.add_heading("Ключевые фрагменты раздела", level=2)
            for item in keys[:8]:
                doc.add_heading(item.rel, level=3)
                doc.add_paragraph(item.role)
                doc.add_paragraph(item.comment)
                add_code_block(doc, item.snippet)

    doc.add_page_break()
    doc.add_heading("Приложение Полный указатель исходников", level=1)
    doc.add_paragraph(
        "В таблице ниже перечислены все включённые исходники. Низкая доля комментариев не означает ошибку; "
        "это указание, где внешняя поясняющая документация особенно полезна."
    )
    rows = []
    for item in sources:
        rows.append([
            item.section,
            item.rel,
            str(item.lines),
            str(item.comment_lines),
            "да" if item.has_intro_comment else "нет",
        ])
    add_table(doc, ["Раздел", "Файл", "Строк", "Комментариев", "Вводный комментарий"], rows, [3.4, 7.8, 1.1, 1.6, 2.0])

    doc.add_heading("Приложение Рекомендации по сопровождению", level=1)
    for text in [
        "Новые алгоритмы добавлять в Algs или Core и подключать к UI через модель тегов, а не через прямые ссылки на формы.",
        "Для устройств сохранять явный lifecycle: поиск, проверка связи, подключение, конфигурация, preview/record, остановка, отключение.",
        "Настройки хранить в проектных и app-конфигурациях через специализированные сервисы, избегая скрытого состояния в визуальных контролах.",
        "Любой модуль, выполняющийся в runtime-цикле, должен иметь предсказуемые буферы и не создавать лишние объекты на каждом тике.",
        "Формы и LFM-файлы лучше документировать через назначение пользовательского сценария: что открывает диалог, какие данные он читает и что применяет по OK.",
    ]:
        doc.add_paragraph(text)

    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    build_doc()
